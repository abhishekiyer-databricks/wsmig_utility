"""Golden-baseline diff for PLAN 16 live QA (master plan, "Golden baseline diff").

Reduces one run's reports to a stable snapshot, then diffs two snapshots:
  • inventory.xlsx      Summary            → count per component
  • export_status.xlsx  Export Summary     → per asset type × status counts
  • import_status.xlsx  Import Summary     → outcome roll-up + per asset type × status counts
                        every per-type tab → natural_key → Import Status (per-object, the full bed)
                        Object Permissions → (object type, import status) counts
                        ACL Parity         → verdict counts

    python3 tests/golden_diff.py snapshot <run_dir>                      # writes <run_dir>/golden_snapshot.json
    python3 tests/golden_diff.py diff <golden_run_dir> <run_dir> [--expect expected.json]

<run_dir> holds reports/{inventory,export_status,import_status}.xlsx (as saved under
~/Desktop/wsmig_runs/plan16_<golden|n>/runN/). `--expect` is the wave's declared expected diff:
{"<section>": {"<key>": "<free-text reason>"}} — a listed key may differ; anything else is a
regression. Exit code 1 when an unexpected difference is found.
"""
from __future__ import annotations

import json
import os
import sys

import openpyxl

_IMPORT_NON_TYPE_TABS = {"Import Summary", "Outstanding", "Object Permissions (ACLs)", "ACL Parity"}


def _rows(wb, name):
    return [r for r in wb[name].iter_rows(values_only=True)] if name in wb.sheetnames else []


def _type_table(rows, first_header):
    """The per-asset-type table that starts at the row whose first two cells are
    (first_header, 'Total'); ends at the first blank row."""
    out, hdr = {}, None
    for r in rows:
        if hdr is None:
            if r and r[0] == first_header and len(r) > 1 and r[1] == "Total":
                hdr = [str(c) for c in r if c is not None]
            continue
        if not r or r[0] is None or str(r[0]).startswith("Account-level"):
            break
        out[str(r[0])] = {h: r[i] for i, h in enumerate(hdr[1:], 1) if r[i]}
    return out


def snapshot(run_dir: str) -> dict:
    rep = os.path.join(run_dir, "reports")
    snap: dict = {}

    wb = openpyxl.load_workbook(os.path.join(rep, "inventory.xlsx"), read_only=True)
    snap["inventory_summary"] = {str(r[0]): r[1] for r in _rows(wb, "Summary")
                                 if r and r[0] and isinstance(r[1], int)}

    wb = openpyxl.load_workbook(os.path.join(rep, "export_status.xlsx"), read_only=True)
    snap["export_by_type"] = _type_table(_rows(wb, "Export Summary"), "Asset Type")

    wb = openpyxl.load_workbook(os.path.join(rep, "import_status.xlsx"), read_only=True)
    summ = _rows(wb, "Import Summary")
    for i, r in enumerate(summ):
        if r and r[0] == "Outcome roll-up":
            hdr, vals = summ[i + 1], summ[i + 2]
            snap["import_rollup"] = {str(h): vals[j] for j, h in enumerate(hdr) if h}
            break
    snap["import_by_type"] = _type_table(summ, "Asset Type")

    per_object: dict = {}
    for name in wb.sheetnames:
        if name in _IMPORT_NON_TYPE_TABS:
            continue
        rows = _rows(wb, name)
        if not rows or rows[0][:3] != ("Asset Type", "Natural Key", "Import Status"):
            continue
        for r in rows[1:]:
            if r and r[1]:
                per_object[f"{r[0]}|{r[1]}"] = str(r[2])
    snap["import_per_object"] = per_object

    acl: dict = {}
    for r in _rows(wb, "Object Permissions (ACLs)")[1:]:
        if r and r[0]:
            k = f"{r[0]}|{r[5]}"
            acl[k] = acl.get(k, 0) + 1
    snap["acl_grants"] = acl

    parity: dict = {}
    for r in _rows(wb, "ACL Parity")[2:]:
        if r and r[2]:
            k = f"{r[0]}|{r[2]}"
            parity[k] = parity.get(k, 0) + 1
    snap["acl_parity"] = parity
    return snap


def _load(path: str) -> dict:
    if os.path.isdir(path):
        cached = os.path.join(path, "golden_snapshot.json")
        return json.load(open(cached)) if os.path.isfile(cached) else snapshot(path)
    return json.load(open(path))


def diff(golden: dict, current: dict, expect: dict) -> list[str]:
    problems = []
    for section in sorted(set(golden) | set(current)):
        g, c = golden.get(section, {}), current.get(section, {})
        for key in sorted(set(g) | set(c), key=str):
            if g.get(key) == c.get(key):
                continue
            reason = (expect.get(section) or {}).get(key)
            tag = f"EXPECTED ({reason})" if reason else "UNEXPECTED"
            problems.append(f"[{tag}] {section} :: {key}: golden={g.get(key)!r} now={c.get(key)!r}")
    return problems


def main(argv):
    if len(argv) >= 2 and argv[0] == "snapshot":
        snap = snapshot(argv[1])
        out = os.path.join(argv[1], "golden_snapshot.json")
        with open(out, "w") as f:
            json.dump(snap, f, indent=1, sort_keys=True, default=str)
        print(f"snapshot written: {out} ({len(snap['import_per_object'])} objects)")
        return 0
    if len(argv) >= 3 and argv[0] == "diff":
        expect = {}
        if "--expect" in argv:
            expect = json.load(open(argv[argv.index("--expect") + 1]))
        lines = diff(_load(argv[1]), _load(argv[2]), expect)
        for line in lines:
            print(line)
        bad = [x for x in lines if x.startswith("[UNEXPECTED]")]
        print(f"\n{len(lines)} differences, {len(bad)} unexpected")
        return 1 if bad else 0
    print(__doc__)
    return 2


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
